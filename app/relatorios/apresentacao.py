"""
Gerador de relatório em PowerPoint (python-pptx) no padrão visual do painel.

Slides 16:9, fundo branco, títulos em Cambria e corpo em Calibri (fontes que vêm com o
Office), azul da marca (#1B365D) na capa e nos destaques. Gráficos são NATIVOS do
PowerPoint (editáveis, com os dados embutidos) — não imagens.

Uso:
    ap = Apresentacao("Estoque", "Posição em 13/10/2026 · todos os segmentos", fonte="relatório Sienge (RPA)")
    ap.capa()
    ap.kpis("Visão geral", [("Valor em estoque", "R$ 41,0 mi", "15.489 posições"), ...])
    ap.barras("Valor por segmento", ["MCMV", ...], [16.6e6, ...], formato="brl")
    ap.tabela("Maiores itens", df, colunas)          # colunas no formato de app.exportacao
    ap.notas("Como ler", ["...", "..."])
    conteudo = ap.salvar()                           # bytes do .pptx
"""
from __future__ import annotations

import io
import math
from datetime import datetime

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.exportacao import NUMERICOS, formatar

AZUL = RGBColor(0x1B, 0x36, 0x5D)
AZUL_CLARO = RGBColor(0xCA, 0xDC, 0xFC)
TEXTO = RGBColor(0x14, 0x21, 0x3D)
TEXTO_2 = RGBColor(0x4A, 0x55, 0x68)
TEXTO_3 = RGBColor(0x5F, 0x6E, 0x85)
CARTAO = RGBColor(0xF1, 0xF4, 0xF8)
LINHA = RGBColor(0xD9, 0xE1, 0xEA)
ZEBRA = RGBColor(0xF6, 0xF8, 0xFB)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
VERMELHO = RGBColor(0xB4, 0x23, 0x18)
VERDE = RGBColor(0x0F, 0x7B, 0x6C)
SERIES = [RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xEB, 0x68, 0x34), RGBColor(0x1B, 0xAF, 0x7A),
          RGBColor(0xED, 0xA1, 0x00)]
TITULO_FONTE, CORPO_FONTE = "Cambria", "Calibri"

LARGURA, ALTURA = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.6)
TOPO_CONTEUDO = Inches(1.45)

FORMATO_EIXO = {"brl": '"R$" #,##0', "brl4": '"R$" #,##0.00', "num": "#,##0", "num2": "#,##0.00",
                "int": "#,##0", "pct": '0.0"%"'}


def _rgb_hex(c: RGBColor) -> str:
    return str(c)


class Apresentacao:
    def __init__(self, titulo: str, subtitulo: str = "", fonte: str = "", area: str = "Suprimentos"):
        self.titulo, self.subtitulo, self.fonte, self.area = titulo, subtitulo, fonte, area
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = LARGURA, ALTURA
        self.prs.core_properties.title = f"{area} · {titulo}"
        self.prs.core_properties.author = "Painel de Suprimentos"
        self._branco = self.prs.slide_layouts[6]
        self._n = 0
        self.gerado = datetime.now()

    # ── básicos ──────────────────────────────────────────────────────────────
    def _texto(self, slide, x, y, w, h, texto, tamanho=14, cor=TEXTO, negrito=False, fonte=CORPO_FONTE,
               alinhar=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP, italico=False, nome=None):
        caixa = slide.shapes.add_textbox(x, y, w, h)
        if nome:
            caixa.name = nome
        tf = caixa.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ancora
        linhas = texto if isinstance(texto, list) else [texto]
        for i, t in enumerate(linhas):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = alinhar
            r = p.add_run()
            r.text = str(t)
            f = r.font
            f.size, f.bold, f.italic, f.name = Pt(tamanho), negrito, italico, fonte
            f.color.rgb = cor
        return caixa

    def _slide(self, titulo: str | None = None):
        s = self.prs.slides.add_slide(self._branco)
        self._n += 1
        if titulo:
            self._texto(s, MARGEM, Inches(0.45), LARGURA - 2 * MARGEM, Inches(0.75), titulo, 28, TEXTO, True,
                        TITULO_FONTE, nome="Título")
        rodape = f"{self.area} · {self.titulo}" + (f" · fonte: {self.fonte}" if self.fonte else "") + \
                 f" · gerado em {self.gerado:%d/%m/%Y %H:%M}"
        self._texto(s, MARGEM, ALTURA - Inches(0.45), Inches(10.5), Inches(0.25), rodape, 9, TEXTO_3, nome="Rodapé")
        self._texto(s, LARGURA - MARGEM - Inches(1), ALTURA - Inches(0.45), Inches(1), Inches(0.25), str(self._n), 9,
                    TEXTO_3, alinhar=PP_ALIGN.RIGHT, nome="Número")
        return s

    def _nota(self, s, texto, y=None):
        if not texto:
            return
        y = y or ALTURA - Inches(1.05)
        self._texto(s, MARGEM, y, LARGURA - 2 * MARGEM, Inches(0.5), texto, 11, TEXTO_2, italico=True, nome="Nota")

    # ── slides ───────────────────────────────────────────────────────────────
    def capa(self, destaques: list[tuple[str, str]] | None = None):
        s = self.prs.slides.add_slide(self._branco)
        self._n += 1
        fundo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, LARGURA, ALTURA)
        fundo.fill.solid()
        fundo.fill.fore_color.rgb = AZUL
        fundo.line.fill.background()
        fundo.name = "Fundo"
        self._texto(s, MARGEM + Inches(0.1), Inches(0.7), Inches(8), Inches(0.4), f"TELESIL · {self.area.upper()}", 13,
                    AZUL_CLARO, True, nome="Área")
        self._texto(s, MARGEM + Inches(0.1), Inches(2.3), Inches(11), Inches(1.4), self.titulo, 46, BRANCO, True,
                    TITULO_FONTE, nome="Título")
        if self.subtitulo:
            self._texto(s, MARGEM + Inches(0.1), Inches(3.7), Inches(11), Inches(0.9), self.subtitulo, 18, AZUL_CLARO,
                        nome="Subtítulo")
        if destaques:
            larg = Inches(2.9)
            for i, (rot, val) in enumerate(destaques[:4]):
                x = MARGEM + Inches(0.1) + i * (larg + Inches(0.25))
                self._texto(s, x, Inches(5.1), larg, Inches(0.55), val, 26, BRANCO, True, TITULO_FONTE,
                            nome=f"Destaque {i + 1}")
                self._texto(s, x, Inches(5.7), larg, Inches(0.5), rot, 12, AZUL_CLARO, nome=f"Destaque {i + 1} rótulo")
        self._texto(s, MARGEM + Inches(0.1), ALTURA - Inches(0.8), Inches(11), Inches(0.3),
                    f"Gerado em {self.gerado:%d/%m/%Y %H:%M}" + (f" · fonte: {self.fonte}" if self.fonte else ""),
                    11, AZUL_CLARO, nome="Data")
        return s

    def kpis(self, titulo: str, itens: list[tuple], nota: str | None = None):
        """itens: (rótulo, valor, subtexto[, 'alerta'|'bom'])"""
        s = self._slide(titulo)
        n = len(itens)
        colunas = 3 if n > 4 else max(1, n)
        linhas = math.ceil(n / colunas)
        gap = Inches(0.3)
        larg = (LARGURA - 2 * MARGEM - gap * (colunas - 1)) / colunas
        alt = Inches(2.15) if linhas > 1 else Inches(2.6)
        for i, item in enumerate(itens):
            rot, val, sub = item[0], item[1], item[2] if len(item) > 2 else ""
            tom = item[3] if len(item) > 3 else None
            x = MARGEM + (i % colunas) * (larg + gap)
            y = TOPO_CONTEUDO + (i // colunas) * (alt + gap)
            card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, larg, alt)
            card.adjustments[0] = 0.06
            card.fill.solid()
            card.fill.fore_color.rgb = CARTAO
            card.line.fill.background()
            card.shadow.inherit = False
            card.name = f"Cartão {i + 1}"
            cor = VERMELHO if tom == "alerta" else VERDE if tom == "bom" else AZUL
            pad = Inches(0.3)
            self._texto(s, x + pad, y + Inches(0.25), larg - 2 * pad, Inches(0.35), rot, 13, TEXTO_2, nome=f"KPI {i+1}")
            self._texto(s, x + pad, y + Inches(0.7), larg - 2 * pad, Inches(0.8), val, 34 if len(str(val)) < 14 else 26,
                        cor, True, TITULO_FONTE, nome=f"KPI {i + 1} valor")
            if sub:
                self._texto(s, x + pad, y + alt - Inches(0.75), larg - 2 * pad, Inches(0.55), sub, 12, TEXTO_3,
                            nome=f"KPI {i + 1} sub")
        self._nota(s, nota)
        return s

    def _grafico(self, s, tipo, categorias, series: dict, formato: str, x, y, w, h, cores=None, legenda=None,
                 rotulos=True, cores_pontos=None, eixo=None):
        dados = CategoryChartData()
        dados.categories = [str(c)[:48] for c in categorias]
        for nome, valores in series.items():
            dados.add_series(nome, [None if (v is None or (isinstance(v, float) and math.isnan(v))) else float(v)
                                    for v in valores])
        gf = s.shapes.add_chart(tipo, x, y, w, h, dados)
        gf.name = "Gráfico"
        ch = gf.chart
        ch.font.size, ch.font.name = Pt(11), CORPO_FONTE
        ch.font.color.rgb = TEXTO_2
        varias = len(series) > 1
        ch.has_legend = bool(legenda if legenda is not None else varias)
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(11)
        fmt = FORMATO_EIXO.get(formato, "#,##0")
        maximo = max([abs(float(v)) for vs in series.values() for v in vs
                      if v is not None and not (isinstance(v, float) and math.isnan(v))] or [0])
        if formato == "brl" and maximo >= 1e6:
            fmt = '"R$" #,##0.0,," mi"'
        elif formato == "brl" and maximo >= 1e4:
            fmt = '"R$" #,##0," mil"'
        ch.has_title = False
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = LINHA
        va.format.line.fill.background()
        va.tick_labels.number_format, va.tick_labels.number_format_is_linked = fmt, False
        va.tick_labels.font.size = Pt(10)
        ca = ch.category_axis
        ca.format.line.color.rgb = LINHA
        ca.has_major_gridlines = False
        ca.tick_labels.font.size = Pt(10.5)
        plot = ch.plots[0]
        if hasattr(plot, "gap_width"):
            plot.gap_width = 55
        for i, serie in enumerate(plot.series):
            cor = (cores or SERIES)[i % len(cores or SERIES)]
            if tipo in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                serie.format.line.color.rgb = cor
                serie.format.line.width = Pt(2.5)
                serie.smooth = False
                serie.marker.format.fill.solid()
                serie.marker.format.fill.fore_color.rgb = cor
            else:
                serie.format.fill.solid()
                serie.format.fill.fore_color.rgb = cor
        if cores_pontos and len(plot.series) == 1:
            for i, cor in enumerate(cores_pontos):
                ponto = plot.series[0].points[i]
                ponto.format.fill.solid()
                ponto.format.fill.fore_color.rgb = cor
        if eixo:
            va.minimum_scale, va.maximum_scale = eixo
        if rotulos:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format, dl.number_format_is_linked = fmt, False
            dl.font.size = Pt(10)
            dl.font.color.rgb = TEXTO_2
            if tipo in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif tipo in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                dl.position = XL_LABEL_POSITION.ABOVE
        if tipo == XL_CHART_TYPE.BAR_CLUSTERED:
            ca.reverse_order = True   # maior em cima
            if rotulos:               # rótulo em cada barra: o eixo de valores vira ruído
                va.visible = False
                va.has_major_gridlines = False
        return ch

    def barras(self, titulo: str, categorias, valores, formato="brl", horizontal=True, nota=None, nome_serie="Valor",
               comentario: list[str] | None = None, cor=None, cores_pontos=None):
        s = self._slide(titulo)
        largura = LARGURA - 2 * MARGEM - (Inches(3.6) if comentario else 0)
        n = len(categorias)
        if n == 0:
            self._texto(s, MARGEM, TOPO_CONTEUDO, largura, Inches(1), "Sem dados para os filtros escolhidos.", 16, TEXTO_2)
            return s
        tipo = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        self._grafico(s, tipo, categorias, {nome_serie: valores}, formato, MARGEM, TOPO_CONTEUDO, largura,
                      ALTURA - TOPO_CONTEUDO - Inches(1.2), cores=[cor or AZUL], cores_pontos=cores_pontos)
        if comentario:
            self._comentario(s, comentario)
        self._nota(s, nota)
        return s

    def linhas(self, titulo: str, categorias, series: dict, formato="brl", nota=None, comentario=None, colunas=False,
               eixo: tuple | None = None):
        s = self._slide(titulo)
        largura = LARGURA - 2 * MARGEM - (Inches(3.6) if comentario else 0)
        if not len(categorias):
            self._texto(s, MARGEM, TOPO_CONTEUDO, largura, Inches(1), "Sem dados para os filtros escolhidos.", 16, TEXTO_2)
            return s
        tipo = XL_CHART_TYPE.COLUMN_CLUSTERED if colunas else XL_CHART_TYPE.LINE_MARKERS
        self._grafico(s, tipo, categorias, series, formato, MARGEM, TOPO_CONTEUDO, largura,
                      ALTURA - TOPO_CONTEUDO - Inches(1.2), rotulos=len(categorias) <= 14, eixo=eixo)
        if comentario:
            self._comentario(s, comentario)
        self._nota(s, nota)
        return s

    def _comentario(self, s, itens: list[str]):
        x = LARGURA - MARGEM - Inches(3.3)
        caixa = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, TOPO_CONTEUDO, Inches(3.3),
                                   ALTURA - TOPO_CONTEUDO - Inches(1.2))
        caixa.adjustments[0] = 0.04
        caixa.fill.solid()
        caixa.fill.fore_color.rgb = CARTAO
        caixa.line.fill.background()
        caixa.name = "Comentário"
        tf = caixa.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, t in enumerate(itens):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            r = p.add_run()
            r.text = t
            r.font.size, r.font.name = Pt(13 if i else 14), CORPO_FONTE
            r.font.bold = i == 0
            r.font.color.rgb = TEXTO if i else AZUL

    def tabela(self, titulo: str, df: pd.DataFrame, colunas: list[dict], nota: str | None = None,
               max_linhas: int = 11, total: bool = True):
        s = self._slide(titulo)
        cols = [c for c in colunas if c["campo"] in df]
        d = df.head(max_linhas)
        if d.empty:
            self._texto(s, MARGEM, TOPO_CONTEUDO, Inches(10), Inches(1), "Sem dados para os filtros escolhidos.", 16,
                        TEXTO_2)
            return s
        tot = {}
        if total:
            for c in cols:
                if c.get("total"):
                    tot[c["campo"]] = float(pd.to_numeric(df[c["campo"]], errors="coerce").sum())
        n_linhas = len(d) + 1 + (1 if tot else 0)
        alt_linha = Inches(0.34)
        largura = LARGURA - 2 * MARGEM
        padrao = {"texto": 2.2, "brl": 1.5, "brl4": 1.3}
        pesos = [max(c.get("peso", 0) or 0, padrao[c["formato"]]) if c.get("formato") in ("brl", "brl4")
                 else c.get("peso", padrao.get(c.get("formato", "texto"), 1.0)) for c in cols]
        tab = s.shapes.add_table(n_linhas, len(cols), MARGEM, TOPO_CONTEUDO, largura, alt_linha * n_linhas).table
        larguras_pol = []
        for j, p in enumerate(pesos):
            tab.columns[j].width = Emu(int(largura * p / sum(pesos)))
            larguras_pol.append(tab.columns[j].width / 914400)

        def cel(i, j, texto, negrito=False, cor=TEXTO, fundo=None, tam=11, direita=False):
            c = tab.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if direita else PP_ALIGN.LEFT
            r = p.add_run()
            r.text = texto
            r.font.size, r.font.bold, r.font.name = Pt(tam), negrito, CORPO_FONTE
            r.font.color.rgb = cor
            c.fill.solid()
            c.fill.fore_color.rgb = fundo or BRANCO

        for j, c in enumerate(cols):
            cel(0, j, c["rotulo"], True, BRANCO, AZUL, 10.5, c.get("formato") in NUMERICOS)
        tab.rows[0].height = alt_linha
        for i, linha in enumerate(d[[c["campo"] for c in cols]].itertuples(index=False), start=1):
            for j, (v, c) in enumerate(zip(linha, cols)):
                txt = formatar(v, c.get("formato", "texto"))
                # uma linha por célula: corta TEXTO pela largura da coluna (Calibri 11 ≈ 11 caracteres/pol.,
                # com folga para maiúsculas); números nunca são cortados
                limite = max(6, int(larguras_pol[j] * 11.5)) if c.get("formato", "texto") == "texto" else 99
                cel(i, j, txt if len(txt) <= limite else txt[:limite - 1] + "…", fundo=ZEBRA if i % 2 == 0 else BRANCO,
                    direita=c.get("formato") in NUMERICOS)
        if tot:
            i = n_linhas - 1
            for j, c in enumerate(cols):
                txt = formatar(tot[c["campo"]], c["formato"]) if c["campo"] in tot else ("Total" if j == 0 else "")
                cel(i, j, txt, True, TEXTO, CARTAO, direita=c.get("formato") in NUMERICOS)
        extra = f"Mostrando {len(d)} de {len(df)} linhas (total sobre todas). " if len(df) > len(d) else ""
        self._nota(s, (extra + (nota or "")).strip() or None)
        return s

    def notas(self, titulo: str, itens: list[str]):
        s = self._slide(titulo)
        y = TOPO_CONTEUDO
        for i, t in enumerate(itens):
            bola = s.shapes.add_shape(MSO_SHAPE.OVAL, MARGEM, y + Inches(0.02), Inches(0.42), Inches(0.42))
            bola.fill.solid()
            bola.fill.fore_color.rgb = AZUL
            bola.line.fill.background()
            bola.name = f"Marcador {i + 1}"
            tf = bola.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(i + 1)
            r.font.size, r.font.bold, r.font.name = Pt(13), True, CORPO_FONTE
            r.font.color.rgb = BRANCO
            caixa = self._texto(s, MARGEM + Inches(0.65), y, LARGURA - 2 * MARGEM - Inches(0.65), Inches(0.9), t, 15,
                                TEXTO, nome=f"Nota {i + 1}")
            y += Inches(0.95)
            if y > ALTURA - Inches(1.3):
                break
        return s

    def salvar(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


def nome_arquivo(area: str) -> str:
    import re
    import unicodedata
    base = unicodedata.normalize("NFKD", area.lower()).encode("ascii", "ignore").decode("ascii")
    base = re.sub(r"[^a-z0-9]+", "_", base).strip("_") or "relatorio"
    return f"relatorio_{base}_{datetime.now():%Y%m%d_%H%M}.pptx"
